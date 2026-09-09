from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

# ------------------------------------------------------------
# 1. CREATE A NEW EXCEL FILE
# ------------------------------------------------------------

wb = Workbook()          # Creates a new Excel workbook object
ws = wb.active           # Gets the default active sheet ("Sheet")

ws['A1'] = "Hello World!"   # Writes text into cell A1

wb.save('example.xlsx')     # Saves the workbook to disk


# ------------------------------------------------------------
# 2. LOAD EXISTING EXCEL FILE
# ------------------------------------------------------------

wb = load_workbook('example.xlsx')   # Opens the file we just saved
ws = wb['Sheet']                     # Accesses the sheet by name

print(ws['A1'].value)                # Reads and prints the value in A1


# ------------------------------------------------------------
# 3. READ ALL CELLS IN ROW 1
# ------------------------------------------------------------

wb = load_workbook('example.xlsx')
ws = wb['Sheet']

for cell in ws['1']:                 # ws['1'] = entire first row
    print(cell.value)                # Print each cell's value


# ------------------------------------------------------------
# 4. WRITE NEW DATA
# ------------------------------------------------------------

ws['B2'] = 'New Data'                # Write into cell B2
wb.save('example.xlsx')


# ------------------------------------------------------------
# 5. APPLY FORMATTING + FORMULAS
# ------------------------------------------------------------

ws['A1'].font = Font(bold=True)      # Make A1 bold

ws['C1'] = '=SUM(A1:B1)'             # Excel formula: sum A1 + B1

wb.save('example.xlsx')


# ------------------------------------------------------------
# 6. MERGE + FREEZE PANES
# ------------------------------------------------------------

ws.merge_cells('A1:B1')              # Merge A1 and B1 into one cell
ws.freeze_panes = 'A2'               # Freeze row 1

wb.save('example.xlsx')


"""
Mini Project
Automating monthly report generation from Excel data saves time and reduces errors.
Create a monthly report of sales data and produce it in an Excel file with formatting.
"""
