from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment
import os

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

def create_report(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "SALES_REPORT"
    ws.append(["MONTH", "SALES", "REGION"])
    wb.save(path)

'''
What
Creates a new Excel file with:
a sheet named "SALES_REPORT"
a header row

Why
Every pipeline needs a base structure before inserting data.
This function builds the skeleton of the report.
'''

def add_sales_data(path, data_rows):
    wb = load_workbook(path)
    ws = wb.active

    for row in data_rows:
        ws.append(row)

    wb.save(path)

'''
What
Loads the Excel file and appends multiple rows of sales data.

Why
This is the data ingestion step.

Pipelines always separate:
file creation
data insertion
formatting

This keeps the workflow modular.
'''

def apply_formatting(path):
    wb = load_workbook(path)
    ws = wb.active

    # Bold headers
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")


    # Freeze header row
    ws.freeze_panes = "A2"

    wb.save(path)

'''
What
Makes header row bold

Centers header text

Freezes row 1

Why
Formatting is done after data insertion so the pipeline stays clean.
This step improves readability of the final report.

'''

def run_monthly_report_pipeline():
    excel_path = os.path.join(BASE_DIR, "monthly_sales_report.xlsx")

    create_report(excel_path)

    sales_data = [
        ["JAN", 12000, "WEST"],
        ["FEB", 15000, "EAST"],
        ["MAR", 18000, "SOUTH"],
        ["APR", 21000, "NORTH"],
        ["MAY", 19500, "WEST"],
    ]

    add_sales_data(excel_path, sales_data)
    apply_formatting(excel_path)

'''
What
This is the orchestrator:
defines the output file name
creates the report
inserts data
formats the report

Why
Every pipeline needs a single entry point that runs all steps in order.

'''

if __name__ == "__main__":
    run_monthly_report_pipeline()
